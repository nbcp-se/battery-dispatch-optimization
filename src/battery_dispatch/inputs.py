from datetime import datetime, timedelta
from math import isfinite
from pathlib import Path

from openpyxl import load_workbook

from .model import Battery, Period


def read_battery(path: Path) -> Battery:
    workbook = load_workbook(path, read_only=True, data_only=True)

    try:
        values = {row[0]: row[1] for row in workbook["Data"].iter_rows(values_only=True)
                  if row[0] is not None}
        return Battery(*(float(values[key]) for key in (
            "Max charging rate", "Max discharging rate", "Max storage volume",
            "Battery charging efficiency", "Battery discharging efficiency",
        )))

    finally:
        workbook.close()


def _read_market(sheet, interval_minutes: int, start: datetime | None = None,
                 end: datetime | None = None) -> dict[datetime, float]:
    prices = {}
    epoch = datetime(1970, 1, 1)

    for row_number, (timestamp, price) in enumerate(
        sheet.iter_rows(min_row=2, max_col=2, values_only=True), start=2
    ):
        if timestamp is None and price is None:
            continue  # The hourly sheeet contains trailing formated empty rows
        
        label = f"{sheet.title}, row {row_number}"
        
        if not isinstance(timestamp, datetime) or timestamp.tzinfo is not None:
            raise ValueError(f"Invalid timestamp at {label}")
        
        # Excels hourly timestamps drift by milliseconds. Only repair tiny offsets
        seconds = (timestamp - epoch).total_seconds()
        aligned = epoch + timedelta(seconds=round(seconds / (60 * interval_minutes))
                                    * 60 * interval_minutes)
        
        if start is not None and end is not None and not start <= aligned < end:
            continue
        
        if abs((aligned - timestamp).total_seconds()) > 1:
            raise ValueError(f"Timestamp is off the market interval at {label}")
        
        if isinstance(price, bool) or not isinstance(price, (int, float)) or not isfinite(price):
            raise ValueError(f"Missing or invalid price at {label}")
        
        if aligned in prices:
            raise ValueError(f"Duplicate timestamp at {label}")
        
        prices[aligned] = float(price)
    
    return prices


def read_periods(path: Path, start: datetime, days: int) -> list[Period]:
    if days <= 0:
        raise ValueError("Days must be positive")
    
    if start.tzinfo is not None or (start.minute, start.second, start.microsecond) != (0, 0, 0):
        raise ValueError("Start must be a naive whole hour")
    
    workbook = load_workbook(path, read_only=True, data_only=True)
    
    try:
        end = start + timedelta(days=days)
        market1 = _read_market(workbook["Half-hourly data"], 30, start, end)
        market2 = _read_market(workbook["Hourly data"], 60, start, end)
    finally:
        workbook.close()
    
    periods = []
    
    for period_index in range(days * 48):
        timestamp = start + timedelta(minutes=30 * period_index)
        hour = timestamp.replace(minute=0)
        
        if timestamp not in market1 or hour not in market2:
            raise ValueError(f"Missing market data for {timestamp.isoformat()}")
        periods.append(Period(timestamp, market1[timestamp], market2[hour]))
    
    return periods
